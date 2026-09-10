#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从 Excel/CSV 读取 userId + BVN，VT 后回填 target.user_info.id_number。

输入列名（大小写不敏感）：userId/user_id + bvn/id_number/raw_bvn

流程：
  1. 读入明文 BVN（跳过空值）
  2. 先查 nigeria_backend.vt_token_cache(vt_type=4)，未命中再 POST /v2t
  3. dry-run 默认只预览；加 --apply 才 UPDATE target.user_info

Usage:
  python3 scripts/backfill_user_info_id_number_from_bvn.py \\
    --env ./.env \\
    --input /path/无标题.xls

  python3 scripts/backfill_user_info_id_number_from_bvn.py \\
    --env ./.env \\
    --input /path/无标题.xls \\
    --apply --force
"""
from __future__ import annotations

import argparse
import csv
import importlib.util
import re
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import pymysql

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE / "reconcile"))
import env_util  # noqa: E402

_MYSQL_RETRYABLE = frozenset({2003, 2006, 2013, 2014})
_BVN_RE = re.compile(r"^\d{11}$")


def log(msg: str) -> None:
    print("[{0}] {1}".format(time.strftime("%F %T"), msg), flush=True)


def load_migrate_collection():
    path = HERE / "migrate_collection.py"
    spec = importlib.util.spec_from_file_location("migrate_collection", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def normalize_header(name: str) -> str:
    return re.sub(r"[\s_-]+", "", (name or "").strip().lower())


def normalize_bvn(raw: Any) -> str:
    if raw is None:
        return ""
    if isinstance(raw, float):
        if raw != raw:  # NaN
            return ""
        if raw == int(raw):
            s = str(int(raw))
        else:
            s = str(raw).strip()
    else:
        s = str(raw).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def normalize_user_id(raw: Any) -> Optional[int]:
    if raw is None or raw == "":
        return None
    if isinstance(raw, float):
        if raw != raw:
            return None
        return int(raw)
    s = str(raw).strip()
    if not s:
        return None
    if s.endswith(".0"):
        s = s[:-2]
    return int(s)


def read_rows(path: Path) -> List[Tuple[int, str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)
    if suffix == ".xls":
        return _read_xls(path)
    raise SystemExit("unsupported input format: {0} (use .xls/.xlsx/.csv)".format(suffix))


def _pick_columns(fieldnames: Sequence[str]) -> Tuple[str, str]:
    norm = {normalize_header(f): f for f in fieldnames}
    uid_key = None
    for cand in ("userid", "user_id", "uid"):
        if cand in norm:
            uid_key = norm[cand]
            break
    bvn_key = None
    for cand in ("bvn", "idnumber", "rawbvn", "raw_bvn"):
        if cand in norm:
            bvn_key = norm[cand]
            break
    if not uid_key or not bvn_key:
        raise SystemExit(
            "cannot detect columns from {0}; need userId + bvn, got {1}".format(
                fieldnames, list(fieldnames)
            )
        )
    return uid_key, bvn_key


def _read_csv(path: Path) -> List[Tuple[int, str]]:
    out: List[Tuple[int, str]] = []
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            return out
        uid_key, bvn_key = _pick_columns(reader.fieldnames)
        for row in reader:
            uid = normalize_user_id(row.get(uid_key))
            bvn = normalize_bvn(row.get(bvn_key))
            if uid is None or not bvn:
                continue
            out.append((uid, bvn))
    return out


def _read_xls(path: Path) -> List[Tuple[int, str]]:
    try:
        import xlrd  # type: ignore
    except ImportError as e:
        raise SystemExit("reading .xls requires xlrd: pip install xlrd==1.2.0") from e
    wb = xlrd.open_workbook(str(path))
    sh = wb.sheet_by_index(0)
    if sh.nrows < 2:
        return []
    headers = [str(sh.cell_value(0, c)).strip() for c in range(sh.ncols)]
    uid_key, bvn_key = _pick_columns(headers)
    uid_idx = headers.index(uid_key)
    bvn_idx = headers.index(bvn_key)
    out: List[Tuple[int, str]] = []
    for r in range(1, sh.nrows):
        uid = normalize_user_id(sh.cell_value(r, uid_idx))
        bvn = normalize_bvn(sh.cell_value(r, bvn_idx))
        if uid is None or not bvn:
            continue
        out.append((uid, bvn))
    return out


def _read_xlsx(path: Path) -> List[Tuple[int, str]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as e:
        raise SystemExit("reading .xlsx requires openpyxl: pip install openpyxl") from e
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []
    headers = [("" if h is None else str(h)).strip() for h in header]
    uid_key, bvn_key = _pick_columns(headers)
    uid_idx = headers.index(uid_key)
    bvn_idx = headers.index(bvn_key)
    out: List[Tuple[int, str]] = []
    for row in rows:
        if row is None:
            continue
        uid = normalize_user_id(row[uid_idx] if uid_idx < len(row) else None)
        bvn = normalize_bvn(row[bvn_idx] if bvn_idx < len(row) else None)
        if uid is None or not bvn:
            continue
        out.append((uid, bvn))
    wb.close()
    return out


def dedupe_rows(rows: Sequence[Tuple[int, str]]) -> Tuple[List[Tuple[int, str]], List[str]]:
    """同一 user_id 保留最后一条 BVN；冲突时记录 warning。"""
    by_uid: Dict[int, str] = {}
    warnings: List[str] = []
    for uid, bvn in rows:
        prev = by_uid.get(uid)
        if prev is not None and prev != bvn:
            warnings.append("user_id={0}: duplicate BVN {1} -> {2} (keep last)".format(uid, prev, bvn))
        by_uid[uid] = bvn
    ordered = sorted(by_uid.items(), key=lambda x: x[0])
    return ordered, warnings


@dataclass
class PlanRow:
    user_id: int
    bvn: str
    token: str
    current_id_number: str
    action: str  # update | skip_same | skip_conflict | skip_missing | skip_invalid_bvn


def mysql_retry(op_name: str, fn, *, retries: int = 5):
    last: Optional[BaseException] = None
    for attempt in range(1, retries + 1):
        try:
            return fn()
        except (pymysql.MySQLError, OSError) as e:
            last = e
            code = e.args[0] if getattr(e, "args", None) else None
            if code not in _MYSQL_RETRYABLE and not isinstance(e, OSError):
                raise
            if attempt >= retries:
                break
            sleep_s = min(2 ** attempt, 60)
            log("MySQL {0} on {1}, retry {2}/{3} after {4}s".format(code, op_name, attempt, retries, sleep_s))
            time.sleep(sleep_s)
    assert last is not None
    raise last


def fetch_current_id_numbers(conn, user_ids: Sequence[int], *, chunk: int = 500) -> Dict[int, str]:
    out: Dict[int, str] = {}
    if not user_ids:
        return out

    def _run():
        nonlocal out
        with conn.cursor() as cur:
            for i in range(0, len(user_ids), chunk):
                part = user_ids[i : i + chunk]
                ph = ",".join(["%s"] * len(part))
                cur.execute(
                    "SELECT user_id, IFNULL(id_number, '') AS id_number "
                    "FROM user_info WHERE user_id IN ({0})".format(ph),
                    part,
                )
                for row in cur.fetchall():
                    out[int(row["user_id"])] = (row["id_number"] or "").strip()
        return out

    return mysql_retry("fetch user_info.id_number", _run)


def apply_updates(
    conn,
    plans: Sequence[PlanRow],
    *,
    batch_size: int,
    apply: bool,
) -> Tuple[int, int]:
    todo = [p for p in plans if p.action == "update"]
    if not todo:
        return 0, 0
    if not apply:
        return len(todo), 0

    updated = 0

    def _run_batch(batch: Sequence[PlanRow]) -> int:
        n = 0
        with conn.cursor() as cur:
            for p in batch:
                cur.execute(
                    "UPDATE user_info SET id_number = %s WHERE user_id = %s",
                    (p.token, p.user_id),
                )
                n += cur.rowcount
        conn.commit()
        return n

    for i in range(0, len(todo), batch_size):
        batch = todo[i : i + batch_size]
        updated += mysql_retry("update user_info.id_number", lambda b=batch: _run_batch(b))
        log("updated progress: {0}/{1}".format(min(i + batch_size, len(todo)), len(todo)))
    return len(todo), updated


def build_plans(
    rows: Sequence[Tuple[int, str]],
    tokens: Dict[str, str],
    current: Dict[int, str],
    *,
    force: bool,
) -> Tuple[List[PlanRow], List[str]]:
    plans: List[PlanRow] = []
    warnings: List[str] = []
    for uid, bvn in rows:
        if not _BVN_RE.match(bvn):
            plans.append(
                PlanRow(uid, bvn, "", current.get(uid, ""), "skip_invalid_bvn")
            )
            warnings.append("user_id={0}: invalid BVN format `{1}`".format(uid, bvn))
            continue
        token = (tokens.get(bvn) or "").strip()
        if not token:
            warnings.append("user_id={0}: VT returned empty token for BVN `{1}`".format(uid, bvn))
            continue
        cur = current.get(uid)
        if cur is None:
            plans.append(PlanRow(uid, bvn, token, "", "skip_missing"))
            warnings.append("user_id={0}: user_info row not found".format(uid))
            continue
        if cur == token:
            plans.append(PlanRow(uid, bvn, token, cur, "skip_same"))
            continue
        if cur and not force:
            plans.append(PlanRow(uid, bvn, token, cur, "skip_conflict"))
            warnings.append(
                "user_id={0}: id_number already set (`{1}`), use --force to overwrite".format(uid, cur[:24])
            )
            continue
        plans.append(PlanRow(uid, bvn, token, cur, "update"))
    return plans, warnings


def main() -> None:
    parser = argparse.ArgumentParser(description="VT BVN and backfill target.user_info.id_number")
    parser.add_argument("--env", default=str(HERE.parent / ".env"), help="path to .env")
    parser.add_argument("--input", required=True, help="xls/xlsx/csv with userId + bvn")
    parser.add_argument("--apply", action="store_true", help="execute UPDATE (default dry-run)")
    parser.add_argument(
        "--force",
        action="store_true",
        help="overwrite non-empty id_number when token differs",
    )
    parser.add_argument("--batch-size", type=int, default=200, help="UPDATE batch size")
    parser.add_argument("--vt-batch-size", type=int, default=2000, help="VT /v2t batch size")
    parser.add_argument("--limit", type=int, default=0, help="only process first N rows after dedupe")
    args = parser.parse_args()

    input_path = Path(args.input).expanduser()
    if not input_path.is_file():
        raise SystemExit("input not found: {0}".format(input_path))

    cfg = env_util.load_env(Path(args.env))
    mc = load_migrate_collection()

    raw_rows = read_rows(input_path)
    rows, dedupe_warn = dedupe_rows(raw_rows)
    if args.limit > 0:
        rows = rows[: args.limit]

    log("input={0} raw_rows={1} deduped={2}".format(input_path, len(raw_rows), len(rows)))
    for w in dedupe_warn[:20]:
        log("WARN: " + w)
    if len(dedupe_warn) > 20:
        log("WARN: ... and {0} more dedupe warnings".format(len(dedupe_warn) - 20))

    unique_bvns = sorted({bvn for _, bvn in rows})
    log("unique_bvn={0}".format(len(unique_bvns)))

    vt_url = (
        cfg.get("VT_BASE_URL")
        or cfg.get("VT_URL")
        or mc.DEFAULT_VT_URL
    )
    source_db = mc.DB(
        mc.DbConfig(
            host=cfg["SOURCE_MYSQL_HOST"],
            port=int(cfg.get("SOURCE_MYSQL_PORT") or 3306),
            user=cfg["SOURCE_MYSQL_USER"],
            password=cfg["SOURCE_MYSQL_PASSWORD"],
            database=cfg.get("SOURCE_MYSQL_DATABASE") or "nigeria_backend",
        ),
        readonly=True,
    )
    vt = mc.VtClient(
        vt_url,
        dry_run=not args.apply,
        db=source_db,
        http_batch_size=max(1, args.vt_batch_size),
    )

    pairs = [(vt.VT_ID_NUMBER, bvn) for bvn in unique_bvns]
    log("VT resolving {0} BVNs...".format(len(unique_bvns)))
    tokens = vt.resolve(pairs)
    log("VT done: resolved={0}/{1}".format(len(tokens), len(unique_bvns)))

    target = env_util.connect_target(cfg)
    try:
        user_ids = [uid for uid, _ in rows]
        current = fetch_current_id_numbers(target, user_ids)
        plans, plan_warn = build_plans(rows, tokens, current, force=args.force)

        counts: Dict[str, int] = {}
        for p in plans:
            counts[p.action] = counts.get(p.action, 0) + 1

        log(
            "plan: update={0} skip_same={1} skip_conflict={2} skip_missing={3} skip_invalid_bvn={4}".format(
                counts.get("update", 0),
                counts.get("skip_same", 0),
                counts.get("skip_conflict", 0),
                counts.get("skip_missing", 0),
                counts.get("skip_invalid_bvn", 0),
            )
        )
        for w in plan_warn[:30]:
            log("WARN: " + w)
        if len(plan_warn) > 30:
            log("WARN: ... and {0} more plan warnings".format(len(plan_warn) - 30))

        sample = [p for p in plans if p.action == "update"][:10]
        if sample:
            log("sample updates:")
            for p in sample:
                log(
                    "  user_id={0} bvn={1} current=`{2}` -> token=`{3}`".format(
                        p.user_id,
                        p.bvn,
                        (p.current_id_number or "")[:32],
                        p.token[:32],
                    )
                )

        todo, updated = apply_updates(
            target,
            plans,
            batch_size=max(1, args.batch_size),
            apply=args.apply,
        )
        if args.apply:
            log("done: attempted={0} rowcount={1}".format(todo, updated))
        else:
            log("dry-run only; add --apply to UPDATE {0} rows".format(todo))
    finally:
        env_util.close_conn(target)


if __name__ == "__main__":
    main()
